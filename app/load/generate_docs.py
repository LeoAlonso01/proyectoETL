import datetime
import logging
import os
from pathlib import Path

import pandas as pd
from docx import Document
from docx2pdf import convert

log = logging.getLogger("load.generate_docs")


def dated_folder(base_output_dir: str) -> str:
    fecha_actual = datetime.datetime.now()
    nombre_carpeta = fecha_actual.strftime("%d_%m_%Y")
    ruta = os.path.join(base_output_dir, nombre_carpeta)
    Path(ruta).mkdir(parents=True, exist_ok=True)
    return ruta


def check_value(val, positivo=(False, "SI", "PERSONA_MORAL")):
    if val in positivo:
        return "X", ""
    return "", "X"


def generate_docs_from_csv(csv_path: str, template_docx_path: str, output_base_dir: str) -> None:
    df = pd.read_csv(csv_path)

    # Filtrado como tu script
    if "encabezado_tipoDeclaracion" in df.columns:
        df = df[df["encabezado_tipoDeclaracion"] != "NOTA"].reset_index(drop=True)

    out_dir = dated_folder(output_base_dir)

    for _, row in df.iterrows():
        doc = Document(template_docx_path)

        bancoSi, bancoNo = check_value(row.get("declaracion_inversionesCuentasValores_ninguno"))
        bienesSi, bienesNo = check_value(row.get("declaracion_bienesInmuebles_ninguno"))
        vehiculosSi, vehiculosNo = check_value(row.get("declaracion_vehiculos_ninguno"))
        extranjeroSi, extranjeroNo = check_value(row.get("domicilioExtranjero"))
        moralSi, moralNo = check_value(row.get("moral_o_fisica"))
        ingresHogarSi, ingresHogarNo = check_value(row.get("declaracion_ingresos_ingresoNetoParejaDependiente_remuneracion_monto"))
        otroSi, otroNo = check_value(row.get("declaracion_ingresos_otrosIngresosTotal_monto"))
        dependientSi, dependientNo = check_value(row.get("declaracion_datosDependientesEconomicos_ninguno"))

        reemplazos_parrafos = {
            "{Tipo}": str(row.get("encabezado_tipoDeclaracion", "")),
            "{Fecha}": str(row.get("encabezado_fechaActualizacion", "")),
            "{Folio}": str(row.get("_id", "")),
        }

        reemplazos_tablas = {
            "{Nombre}": str(row.get("nombre", "")),
            "{Area}": str(row.get("areaAdscripcion", "")),
            "{Empleo}": str(row.get("empleoCargoComision", "")),
            "{FechaIngreso}": str(row.get("fechaEncargo", "")),
            "{Grado}": str(row.get("max_nivel_academico", "")),
            "{BancoSi}": bancoSi,
            "{BancoNo}": bancoNo,
            "{BienesSi}": bienesSi,
            "{BienesNo}": bienesNo,
            "{VehiculosSi}": vehiculosSi,
            "{VehiculosNo}": vehiculosNo,
            "{ExtranjeroSi}": extranjeroSi,
            "{ExtranjeroNo}": extranjeroNo,
            "{MoralSi}": moralSi,
            "{MoralNo}": moralNo,
            "{IngresHogarSi}": ingresHogarSi,
            "{IngresHogarNo}": ingresHogarNo,
            "{OtroSi}": otroSi,
            "{OtroNo}": otroNo,
            "{DependientSi}": dependientSi,
            "{DependientNo}": dependientNo,
        }

        # Párrafos
        for paragraph in doc.paragraphs:
            for marcador, reemplazo in reemplazos_parrafos.items():
                if marcador in paragraph.text:
                    for run in paragraph.runs:
                        run.text = run.text.replace(marcador, reemplazo)

        # Tablas
        for table in doc.tables:
            for table_row in table.rows:
                for cell in table_row.cells:
                    for marcador, reemplazo in reemplazos_tablas.items():
                        if marcador in cell.text:
                            cell.text = cell.text.replace(marcador, reemplazo)

        safe_name = str(row.get("nombre", "SIN_NOMBRE")).replace("/", "_").replace("\\", "_")
        tipo = str(row.get("encabezado_tipoDeclaracion", "DECLARACION"))

        ruta_docx = os.path.join(out_dir, f"{safe_name}_{tipo}.docx")
        ruta_pdf = os.path.join(out_dir, f"{safe_name}_{tipo}.pdf")

        doc.save(ruta_docx)
        try:
            convert(ruta_docx, ruta_pdf)
        except Exception as e:
            log.error("PDF conversion failed for %s: %s", ruta_docx, e)

    log.info("Documents generated in: %s", out_dir)
